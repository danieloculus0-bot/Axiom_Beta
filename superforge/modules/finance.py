from __future__ import annotations

import ast
import csv
import io
import json
import re
import uuid
from datetime import date
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from typing import Any

from openpyxl import Workbook, load_workbook

from ..audit import record_event
from ..db import db
from ..event_bus import publish

_CENT = Decimal("0.01")
_ALLOWED_ACCOUNT_TYPES = {"asset", "liability", "equity", "revenue", "expense"}
_ALLOWED_PAY_TYPES = {"hourly", "salary"}
_SAFE_FUNCS = {
    "ABS": abs,
    "MIN": min,
    "MAX": max,
    "ROUND": round,
    "SUM": lambda *args: sum(args),
}

DEFAULT_ACCOUNTS = [
    ("1000", "Operating Cash", "asset"),
    ("1200", "Accounts Receivable", "asset"),
    ("2000", "Accounts Payable", "liability"),
    ("2100", "Payroll Tax Payable", "liability"),
    ("2110", "Payroll Deductions Payable", "liability"),
    ("2120", "Payroll Clearing", "liability"),
    ("3000", "Owner Equity", "equity"),
    ("4000", "Sales / Revenue", "revenue"),
    ("5000", "Cost of Goods Sold", "expense"),
    ("6100", "Wages & Salaries", "expense"),
    ("6110", "Employer Payroll Taxes", "expense"),
    ("6200", "Benefits & Other Payroll Expense", "expense"),
    ("7000", "General & Administrative Expense", "expense"),
]


def _d(value: Any) -> Decimal:
    try:
        return Decimal(str(0 if value in (None, "") else value))
    except (InvalidOperation, ValueError, TypeError):
        raise ValueError(f"invalid numeric value: {value!r}")


def money(value: Any) -> float:
    return float(_d(value).quantize(_CENT, rounding=ROUND_HALF_UP))


def seed_finance_defaults() -> None:
    with db() as con:
        for code, name, account_type in DEFAULT_ACCOUNTS:
            con.execute(
                """INSERT INTO finance_accounts(code,name,account_type,active)
                   VALUES(?,?,?,1)
                   ON CONFLICT(code) DO NOTHING""",
                (code, name, account_type),
            )


def create_account(code: str, name: str, account_type: str, *, actor: str = "local") -> int:
    code = (code or "").strip()
    name = (name or "").strip()
    account_type = (account_type or "").strip().lower()
    if not code or not name:
        raise ValueError("account code and name are required")
    if account_type not in _ALLOWED_ACCOUNT_TYPES:
        raise ValueError("invalid account type")
    with db() as con:
        cur = con.execute(
            "INSERT INTO finance_accounts(code,name,account_type,active) VALUES(?,?,?,1)",
            (code, name, account_type),
        )
        account_id = int(cur.lastrowid)
    record_event(
        event_type="FINANCE_ACCOUNT",
        action="CREATE",
        module="accounting",
        entity_type="finance_account",
        entity_id=account_id,
        actor=actor,
        data={"code": code, "name": name, "account_type": account_type},
    )
    return account_id


def _account_id(con, code_or_id: Any) -> int:
    if isinstance(code_or_id, int):
        row = con.execute("SELECT id FROM finance_accounts WHERE id=?", (code_or_id,)).fetchone()
    else:
        value = str(code_or_id or "").strip()
        row = con.execute("SELECT id FROM finance_accounts WHERE code=?", (value,)).fetchone()
        if not row and value.isdigit():
            row = con.execute("SELECT id FROM finance_accounts WHERE id=?", (int(value),)).fetchone()
    if not row:
        raise ValueError(f"account not found: {code_or_id}")
    return int(row["id"])


def post_journal_entry(
    entry_date: str,
    memo: str,
    lines: list[dict[str, Any]],
    *,
    actor: str = "local",
    source_module: str = "accounting",
    source_entity_type: str = "",
    source_entity_id: str = "",
    reference: str = "",
) -> int:
    if not lines or len(lines) < 2:
        raise ValueError("journal entry requires at least two lines")
    normalized: list[dict[str, Any]] = []
    debit_total = Decimal("0")
    credit_total = Decimal("0")
    with db() as con:
        for raw in lines:
            debit = _d(raw.get("debit"))
            credit = _d(raw.get("credit"))
            if debit < 0 or credit < 0:
                raise ValueError("debits and credits cannot be negative")
            if debit and credit:
                raise ValueError("one journal line cannot contain both a debit and a credit")
            if not debit and not credit:
                continue
            account_id = _account_id(con, raw.get("account_id") or raw.get("account_code"))
            normalized.append(
                {
                    "account_id": account_id,
                    "description": str(raw.get("description") or ""),
                    "debit": money(debit),
                    "credit": money(credit),
                    "employee_ref": str(raw.get("employee_ref") or ""),
                    "job_ref": str(raw.get("job_ref") or ""),
                }
            )
            debit_total += debit
            credit_total += credit
        if len(normalized) < 2:
            raise ValueError("journal entry requires at least two non-zero lines")
        if debit_total.quantize(_CENT) != credit_total.quantize(_CENT):
            raise ValueError(
                f"journal entry is not balanced: debits={money(debit_total):.2f} credits={money(credit_total):.2f}"
            )
        journal_number = f"JE-{date.today().year}-{uuid.uuid4().hex[:8].upper()}"
        cur = con.execute(
            """INSERT INTO finance_journals(
                 journal_number,entry_date,memo,reference,source_module,source_entity_type,
                 source_entity_id,status,posted_by,posted_at
               ) VALUES(?,?,?,?,?,?,?,'posted',?,CURRENT_TIMESTAMP)""",
            (
                journal_number,
                entry_date or date.today().isoformat(),
                memo or "",
                reference or "",
                source_module,
                source_entity_type,
                str(source_entity_id or ""),
                actor,
            ),
        )
        journal_id = int(cur.lastrowid)
        for line_no, line in enumerate(normalized, start=1):
            con.execute(
                """INSERT INTO finance_journal_lines(
                     journal_id,line_no,account_id,description,debit,credit,employee_ref,job_ref
                   ) VALUES(?,?,?,?,?,?,?,?)""",
                (
                    journal_id,
                    line_no,
                    line["account_id"],
                    line["description"],
                    line["debit"],
                    line["credit"],
                    line["employee_ref"],
                    line["job_ref"],
                ),
            )
    payload = {
        "journal_number": journal_number,
        "entry_date": entry_date,
        "debits": money(debit_total),
        "credits": money(credit_total),
        "reference": reference,
    }
    record_event(
        event_type="FINANCE_JOURNAL",
        action="POST",
        module="accounting",
        source_module=source_module,
        target_module="accounting",
        entity_type="finance_journal",
        entity_id=journal_id,
        actor=actor,
        reason=memo,
        data=payload,
    )
    publish(
        "finance.journal.posted",
        source_module="accounting",
        entity_type="finance_journal",
        entity_id=str(journal_id),
        actor=actor,
        reason=memo,
        payload=payload,
    )
    return journal_id


def create_employee(
    employee_ref: str,
    display_name: str,
    *,
    department: str = "",
    pay_type: str = "hourly",
    hourly_rate: Any = 0,
    annual_salary: Any = 0,
    pay_periods_per_year: int = 52,
    overtime_multiplier: Any = 1.5,
    default_payment_method: str = "check",
    actor: str = "local",
) -> int:
    employee_ref = (employee_ref or "").strip()
    display_name = (display_name or "").strip()
    pay_type = (pay_type or "hourly").strip().lower()
    if not employee_ref or not display_name:
        raise ValueError("employee reference and display name are required")
    if pay_type not in _ALLOWED_PAY_TYPES:
        raise ValueError("pay type must be hourly or salary")
    periods = max(1, int(pay_periods_per_year or 52))
    with db() as con:
        con.execute(
            """INSERT INTO payroll_employees(
                 employee_ref,display_name,department,pay_type,hourly_rate,annual_salary,
                 pay_periods_per_year,overtime_multiplier,default_payment_method,status
               ) VALUES(?,?,?,?,?,?,?,?,?,'active')
               ON CONFLICT(employee_ref) DO UPDATE SET
                 display_name=excluded.display_name,department=excluded.department,
                 pay_type=excluded.pay_type,hourly_rate=excluded.hourly_rate,
                 annual_salary=excluded.annual_salary,pay_periods_per_year=excluded.pay_periods_per_year,
                 overtime_multiplier=excluded.overtime_multiplier,
                 default_payment_method=excluded.default_payment_method,updated_at=CURRENT_TIMESTAMP""",
            (
                employee_ref,
                display_name,
                department,
                pay_type,
                money(hourly_rate),
                money(annual_salary),
                periods,
                float(_d(overtime_multiplier)),
                default_payment_method or "check",
            ),
        )
        employee_id = int(
            con.execute("SELECT id FROM payroll_employees WHERE employee_ref=?", (employee_ref,)).fetchone()["id"]
        )
    record_event(
        event_type="PAYROLL_EMPLOYEE",
        action="UPSERT",
        module="payroll",
        entity_type="payroll_employee",
        entity_id=employee_id,
        actor=actor,
        data={"employee_ref": employee_ref, "department": department, "pay_type": pay_type},
    )
    return employee_id


def create_pay_run(
    period_start: str,
    period_end: str,
    check_date: str,
    *,
    pay_group: str = "WEEKLY",
    run_key: str = "",
    actor: str = "local",
) -> int:
    run_key = (run_key or f"PAY-{period_end}-{uuid.uuid4().hex[:6].upper()}").strip()
    with db() as con:
        cur = con.execute(
            """INSERT INTO payroll_runs(
                 run_key,period_start,period_end,check_date,pay_group,status,created_by
               ) VALUES(?,?,?,?,?,'draft',?)""",
            (run_key, period_start, period_end, check_date, pay_group or "WEEKLY", actor),
        )
        run_id = int(cur.lastrowid)
    record_event(
        event_type="PAYROLL_RUN",
        action="CREATE",
        module="payroll",
        entity_type="payroll_run",
        entity_id=run_id,
        actor=actor,
        data={"run_key": run_key, "period_start": period_start, "period_end": period_end, "check_date": check_date},
    )
    return run_id


def _gross_for_employee(employee: dict[str, Any], regular_hours: Any, overtime_hours: Any, bonus: Any) -> tuple[float, float, float]:
    reg_h = _d(regular_hours)
    ot_h = _d(overtime_hours)
    bonus_d = _d(bonus)
    if reg_h < 0 or ot_h < 0 or bonus_d < 0:
        raise ValueError("hours and bonus cannot be negative")
    if employee["pay_type"] == "salary":
        regular_pay = _d(employee["annual_salary"]) / Decimal(int(employee["pay_periods_per_year"] or 52))
        overtime_pay = Decimal("0")
    else:
        rate = _d(employee["hourly_rate"])
        regular_pay = reg_h * rate
        overtime_pay = ot_h * rate * _d(employee["overtime_multiplier"] or 1.5)
    gross = regular_pay + overtime_pay + bonus_d
    return money(regular_pay), money(overtime_pay), money(gross)


def upsert_pay_item(
    run_id: int,
    employee_ref: str,
    *,
    regular_hours: Any = 0,
    overtime_hours: Any = 0,
    bonus: Any = 0,
    tax_withheld: Any = 0,
    other_deductions: Any = 0,
    payment_method: str = "",
    notes: str = "",
    actor: str = "local",
) -> int:
    with db() as con:
        run = con.execute("SELECT status FROM payroll_runs WHERE id=?", (int(run_id),)).fetchone()
        if not run:
            raise ValueError("pay run not found")
        if run["status"] != "draft":
            raise ValueError("only draft pay runs can be edited")
        row = con.execute("SELECT * FROM payroll_employees WHERE employee_ref=?", (employee_ref,)).fetchone()
        if not row or row["status"] != "active":
            raise ValueError(f"active payroll employee not found: {employee_ref}")
        employee = dict(row)
        regular_pay, overtime_pay, gross = _gross_for_employee(employee, regular_hours, overtime_hours, bonus)
        tax = money(tax_withheld)
        deductions = money(other_deductions)
        if tax < 0 or deductions < 0:
            raise ValueError("withholdings and deductions cannot be negative")
        net = money(_d(gross) - _d(tax) - _d(deductions))
        if net < 0:
            raise ValueError("net pay cannot be negative")
        con.execute(
            """INSERT INTO payroll_items(
                 run_id,employee_id,regular_hours,overtime_hours,regular_pay,overtime_pay,bonus,
                 gross_pay,tax_withheld,other_deductions,net_pay,payment_method,notes,status
               ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,'draft')
               ON CONFLICT(run_id,employee_id) DO UPDATE SET
                 regular_hours=excluded.regular_hours,overtime_hours=excluded.overtime_hours,
                 regular_pay=excluded.regular_pay,overtime_pay=excluded.overtime_pay,bonus=excluded.bonus,
                 gross_pay=excluded.gross_pay,tax_withheld=excluded.tax_withheld,
                 other_deductions=excluded.other_deductions,net_pay=excluded.net_pay,
                 payment_method=excluded.payment_method,notes=excluded.notes,status='draft',
                 updated_at=CURRENT_TIMESTAMP""",
            (
                int(run_id),
                int(employee["id"]),
                float(_d(regular_hours)),
                float(_d(overtime_hours)),
                regular_pay,
                overtime_pay,
                money(bonus),
                gross,
                tax,
                deductions,
                net,
                payment_method or employee["default_payment_method"] or "check",
                notes or "",
            ),
        )
        item_id = int(
            con.execute(
                "SELECT id FROM payroll_items WHERE run_id=? AND employee_id=?",
                (int(run_id), int(employee["id"])),
            ).fetchone()["id"]
        )
    record_event(
        event_type="PAYROLL_ITEM",
        action="UPSERT",
        module="payroll",
        entity_type="payroll_item",
        entity_id=item_id,
        actor=actor,
        data={
            "run_id": int(run_id),
            "employee_ref": employee_ref,
            "gross_pay": gross,
            "tax_withheld": tax,
            "other_deductions": deductions,
            "net_pay": net,
        },
    )
    return item_id


def payroll_run_snapshot(run_id: int) -> dict[str, Any]:
    with db() as con:
        run = con.execute("SELECT * FROM payroll_runs WHERE id=?", (int(run_id),)).fetchone()
        if not run:
            raise ValueError("pay run not found")
        items = [
            dict(r)
            for r in con.execute(
                """SELECT i.*,e.employee_ref,e.display_name,e.department,e.pay_type
                   FROM payroll_items i JOIN payroll_employees e ON e.id=i.employee_id
                   WHERE i.run_id=? ORDER BY e.display_name,e.employee_ref""",
                (int(run_id),),
            )
        ]
    totals = {
        "gross": money(sum(_d(x["gross_pay"]) for x in items)),
        "tax": money(sum(_d(x["tax_withheld"]) for x in items)),
        "deductions": money(sum(_d(x["other_deductions"]) for x in items)),
        "net": money(sum(_d(x["net_pay"]) for x in items)),
    }
    return {"run": dict(run), "items": items, "totals": totals}


def approve_pay_run(run_id: int, *, actor: str = "local") -> int:
    snap = payroll_run_snapshot(run_id)
    if snap["run"]["status"] == "approved":
        return int(snap["run"]["journal_id"])
    if snap["run"]["status"] != "draft":
        raise ValueError("pay run must be draft to approve")
    if not snap["items"]:
        raise ValueError("pay run has no payroll items")
    totals = snap["totals"]
    lines = [
        {"account_code": "6100", "debit": totals["gross"], "description": f"Gross payroll {snap['run']['run_key']}"},
        {"account_code": "2100", "credit": totals["tax"], "description": "Employee tax withholding"},
        {"account_code": "2110", "credit": totals["deductions"], "description": "Employee deductions"},
        {"account_code": "2120", "credit": totals["net"], "description": "Net payroll clearing"},
    ]
    lines = [x for x in lines if money(x.get("debit") or x.get("credit")) != 0]
    journal_id = post_journal_entry(
        snap["run"]["check_date"],
        f"Payroll approval {snap['run']['run_key']}",
        lines,
        actor=actor,
        source_module="payroll",
        source_entity_type="payroll_run",
        source_entity_id=str(run_id),
        reference=snap["run"]["run_key"],
    )
    with db() as con:
        con.execute(
            """UPDATE payroll_runs
               SET status='approved',approved_by=?,approved_at=CURRENT_TIMESTAMP,
                   gross_total=?,tax_total=?,deduction_total=?,net_total=?,journal_id=?
               WHERE id=?""",
            (actor, totals["gross"], totals["tax"], totals["deductions"], totals["net"], journal_id, int(run_id)),
        )
        con.execute("UPDATE payroll_items SET status='approved' WHERE run_id=?", (int(run_id),))
    record_event(
        event_type="PAYROLL_RUN",
        action="APPROVE",
        module="payroll",
        target_module="accounting",
        entity_type="payroll_run",
        entity_id=run_id,
        actor=actor,
        data={**totals, "journal_id": journal_id},
    )
    publish(
        "payroll.run.approved",
        source_module="payroll",
        entity_type="payroll_run",
        entity_id=str(run_id),
        actor=actor,
        payload={**totals, "journal_id": journal_id},
    )
    return journal_id


def mark_pay_run_paid(
    run_id: int,
    *,
    payment_reference: str,
    actor: str = "local",
    payment_method: str = "mixed",
) -> int:
    snap = payroll_run_snapshot(run_id)
    if snap["run"]["status"] == "paid":
        return int(snap["run"]["payment_journal_id"])
    if snap["run"]["status"] != "approved":
        raise ValueError("pay run must be approved before payment")
    if not (payment_reference or "").strip():
        raise ValueError("payment reference is required")
    net = snap["totals"]["net"]
    journal_id = post_journal_entry(
        snap["run"]["check_date"],
        f"Payroll payment {snap['run']['run_key']}",
        [
            {"account_code": "2120", "debit": net, "description": "Clear net payroll"},
            {"account_code": "1000", "credit": net, "description": "Payroll cash payment"},
        ],
        actor=actor,
        source_module="payroll",
        source_entity_type="payroll_run",
        source_entity_id=str(run_id),
        reference=payment_reference,
    )
    with db() as con:
        con.execute(
            """UPDATE payroll_runs
               SET status='paid',payment_reference=?,payment_method=?,paid_by=?,paid_at=CURRENT_TIMESTAMP,
                   payment_journal_id=? WHERE id=?""",
            (payment_reference, payment_method or "mixed", actor, journal_id, int(run_id)),
        )
        con.execute(
            "UPDATE payroll_items SET status='paid',payment_reference=? WHERE run_id=?",
            (payment_reference, int(run_id)),
        )
    record_event(
        event_type="PAYROLL_RUN",
        action="PAID",
        module="payroll",
        target_module="accounting",
        entity_type="payroll_run",
        entity_id=run_id,
        actor=actor,
        data={"net": net, "payment_reference": payment_reference, "payment_journal_id": journal_id},
    )
    publish(
        "payroll.run.paid",
        source_module="payroll",
        entity_type="payroll_run",
        entity_id=str(run_id),
        actor=actor,
        payload={"net": net, "payment_reference": payment_reference, "payment_journal_id": journal_id},
    )
    return journal_id


def trial_balance() -> list[dict[str, Any]]:
    with db() as con:
        rows = con.execute(
            """SELECT a.id,a.code,a.name,a.account_type,
                      COALESCE(SUM(CASE WHEN j.status='posted' THEN l.debit ELSE 0 END),0) debit,
                      COALESCE(SUM(CASE WHEN j.status='posted' THEN l.credit ELSE 0 END),0) credit
               FROM finance_accounts a
               LEFT JOIN finance_journal_lines l ON l.account_id=a.id
               LEFT JOIN finance_journals j ON j.id=l.journal_id
               WHERE a.active=1
               GROUP BY a.id ORDER BY a.code"""
        ).fetchall()
    out = []
    for row in rows:
        d = dict(row)
        d["balance"] = money(_d(d["debit"]) - _d(d["credit"]))
        out.append(d)
    return out


def _safe_name(value: str) -> str:
    s = re.sub(r"[^A-Za-z0-9_]+", "_", (value or "").strip()).strip("_")
    return s or "col"


class _FormulaEvaluator(ast.NodeVisitor):
    def __init__(self, names: dict[str, Any]):
        self.names = {str(k).upper(): v for k, v in names.items()}

    def visit_Expression(self, node):
        return self.visit(node.body)

    def visit_Constant(self, node):
        if isinstance(node.value, (int, float)):
            return node.value
        raise ValueError("only numeric constants are allowed")

    def visit_Name(self, node):
        key = node.id.upper()
        if key not in self.names:
            raise ValueError(f"unknown column: {node.id}")
        try:
            return float(self.names[key])
        except (TypeError, ValueError):
            return 0.0

    def visit_BinOp(self, node):
        left, right = self.visit(node.left), self.visit(node.right)
        if isinstance(node.op, ast.Add): return left + right
        if isinstance(node.op, ast.Sub): return left - right
        if isinstance(node.op, ast.Mult): return left * right
        if isinstance(node.op, ast.Div): return left / right
        if isinstance(node.op, ast.Mod): return left % right
        raise ValueError("unsupported operator")

    def visit_UnaryOp(self, node):
        value = self.visit(node.operand)
        if isinstance(node.op, ast.UAdd): return value
        if isinstance(node.op, ast.USub): return -value
        raise ValueError("unsupported unary operator")

    def visit_Call(self, node):
        if not isinstance(node.func, ast.Name):
            raise ValueError("unsupported function")
        fn = _SAFE_FUNCS.get(node.func.id.upper())
        if not fn:
            raise ValueError(f"unsupported function: {node.func.id}")
        return fn(*[self.visit(arg) for arg in node.args])

    def generic_visit(self, node):
        raise ValueError(f"unsupported formula syntax: {type(node).__name__}")


def evaluate_formula(formula: str, row: dict[str, Any]) -> float:
    expr = (formula or "").strip()
    if expr.startswith("="):
        expr = expr[1:]
    names = {_safe_name(key): value for key, value in row.items()}
    tree = ast.parse(expr, mode="eval")
    return money(_FormulaEvaluator(names).visit(tree))


def resolve_formula_row(row: dict[str, Any], columns: list[str] | None = None) -> dict[str, Any]:
    resolved = dict(row)
    ordered = list(columns or resolved.keys())
    for col in ordered:
        value = resolved.get(col, "")
        if isinstance(value, str) and value.strip().startswith("="):
            resolved[col] = evaluate_formula(value, resolved)
    return resolved


def create_sheet(name: str, columns: list[str], *, template_type: str = "blank", actor: str = "local") -> int:
    clean_columns = [str(x).strip() for x in columns if str(x).strip()]
    if not clean_columns:
        raise ValueError("at least one sheet column is required")
    with db() as con:
        cur = con.execute(
            """INSERT INTO sheet_books(name,template_type,columns_json,rows_json,status,created_by)
               VALUES(?,?,?,?,'draft',?)""",
            (name or "Untitled Sheet", template_type or "blank", json.dumps(clean_columns), "[]", actor),
        )
        sheet_id = int(cur.lastrowid)
    record_event(
        event_type="SHEET",
        action="CREATE",
        module="sheets",
        entity_type="sheet_book",
        entity_id=sheet_id,
        actor=actor,
        data={"name": name, "template_type": template_type, "columns": clean_columns},
    )
    return sheet_id


def template_sheet(template_type: str, *, name: str = "", actor: str = "local") -> int:
    template_type = (template_type or "blank").lower()
    if template_type == "payroll":
        columns = ["Employee Ref", "Regular Hours", "OT Hours", "Bonus", "Tax Withheld", "Other Deductions", "Gross Preview", "Net Preview"]
        default_name = "Payroll Input"
    elif template_type == "journal":
        columns = ["Date", "Account", "Description", "Debit", "Credit", "Employee Ref", "Job Ref"]
        default_name = "Journal Entry"
    else:
        columns = ["A", "B", "C", "D", "E"]
        default_name = "Blank Sheet"
        template_type = "blank"
    return create_sheet(name or default_name, columns, template_type=template_type, actor=actor)


def sheet_snapshot(sheet_id: int, *, evaluate: bool = True) -> dict[str, Any]:
    with db() as con:
        row = con.execute("SELECT * FROM sheet_books WHERE id=?", (int(sheet_id),)).fetchone()
    if not row:
        raise ValueError("sheet not found")
    result = dict(row)
    result["columns"] = json.loads(result["columns_json"])
    result["rows"] = json.loads(result["rows_json"])
    if evaluate:
        evaluated = []
        for raw in result["rows"]:
            out = dict(raw)
            for col in result["columns"]:
                value = out.get(col, "")
                if isinstance(value, str) and value.strip().startswith("="):
                    try:
                        out[col + " (value)"] = evaluate_formula(value, out)
                    except Exception as exc:
                        out[col + " (value)"] = f"#ERR {exc}"
            evaluated.append(out)
        result["evaluated_rows"] = evaluated
    return result


def save_sheet_rows(sheet_id: int, rows: list[dict[str, Any]], *, actor: str = "local") -> None:
    snap = sheet_snapshot(sheet_id, evaluate=False)
    columns = snap["columns"]
    normalized = [{col: row.get(col, "") for col in columns} for row in rows]
    with db() as con:
        con.execute(
            "UPDATE sheet_books SET rows_json=?,updated_at=CURRENT_TIMESTAMP WHERE id=?",
            (json.dumps(normalized, ensure_ascii=False), int(sheet_id)),
        )
    record_event(
        event_type="SHEET",
        action="SAVE_ROWS",
        module="sheets",
        entity_type="sheet_book",
        entity_id=sheet_id,
        actor=actor,
        data={"row_count": len(normalized)},
    )


def import_sheet_bytes(sheet_id: int, filename: str, content: bytes, *, actor: str = "local") -> None:
    lower = (filename or "").lower()
    if lower.endswith(".csv"):
        reader = csv.DictReader(io.StringIO(content.decode("utf-8-sig")))
        columns = list(reader.fieldnames or [])
        rows = [dict(row) for row in reader]
    elif lower.endswith((".xlsx", ".xlsm")):
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
        ws = wb.active
        values = list(ws.iter_rows(values_only=True))
        wb.close()
        if not values:
            columns, rows = [], []
        else:
            columns = [str(x or "").strip() for x in values[0]]
            rows = [
                {columns[i]: ("" if i >= len(row) or row[i] is None else row[i]) for i in range(len(columns))}
                for row in values[1:]
            ]
    else:
        raise ValueError("only CSV, XLSX and XLSM files are supported")
    if not columns:
        raise ValueError("source file has no header row")
    with db() as con:
        con.execute(
            "UPDATE sheet_books SET columns_json=?,rows_json=?,updated_at=CURRENT_TIMESTAMP WHERE id=?",
            (json.dumps(columns), json.dumps(rows, ensure_ascii=False, default=str), int(sheet_id)),
        )
    record_event(
        event_type="SHEET",
        action="IMPORT",
        module="sheets",
        entity_type="sheet_book",
        entity_id=sheet_id,
        actor=actor,
        data={"filename": filename, "row_count": len(rows), "columns": columns},
    )


def export_sheet_csv(sheet_id: int) -> str:
    snap = sheet_snapshot(sheet_id, evaluate=False)
    stream = io.StringIO()
    writer = csv.DictWriter(stream, fieldnames=snap["columns"], extrasaction="ignore")
    writer.writeheader()
    writer.writerows(snap["rows"])
    return stream.getvalue()


def export_sheet_xlsx(sheet_id: int) -> bytes:
    snap = sheet_snapshot(sheet_id, evaluate=False)
    wb = Workbook()
    ws = wb.active
    ws.title = (snap["name"] or "Sheet")[:31]
    ws.append(snap["columns"])
    for row in snap["rows"]:
        ws.append([row.get(col, "") for col in snap["columns"]])
    stream = io.BytesIO()
    wb.save(stream)
    wb.close()
    return stream.getvalue()


def post_payroll_sheet(sheet_id: int, run_id: int, *, actor: str = "local") -> int:
    snap = sheet_snapshot(sheet_id, evaluate=False)
    count = 0
    for raw_row in snap["rows"]:
        row = resolve_formula_row(raw_row, snap["columns"])
        employee_ref = str(row.get("Employee Ref") or row.get("employee_ref") or "").strip()
        if not employee_ref:
            continue
        upsert_pay_item(
            run_id,
            employee_ref,
            regular_hours=row.get("Regular Hours") or row.get("regular_hours") or 0,
            overtime_hours=row.get("OT Hours") or row.get("Overtime Hours") or row.get("overtime_hours") or 0,
            bonus=row.get("Bonus") or row.get("bonus") or 0,
            tax_withheld=row.get("Tax Withheld") or row.get("tax_withheld") or 0,
            other_deductions=row.get("Other Deductions") or row.get("other_deductions") or 0,
            actor=actor,
        )
        count += 1
    with db() as con:
        con.execute(
            """UPDATE sheet_books SET status='posted',posted_entity_type='payroll_run',
               posted_entity_id=?,updated_at=CURRENT_TIMESTAMP WHERE id=?""",
            (str(run_id), int(sheet_id)),
        )
    record_event(
        event_type="SHEET",
        action="POST_TO_PAYROLL",
        module="sheets",
        target_module="payroll",
        entity_type="sheet_book",
        entity_id=sheet_id,
        actor=actor,
        data={"run_id": int(run_id), "items_posted": count},
    )
    return count


def post_journal_sheet(sheet_id: int, *, actor: str = "local", memo: str = "Sheet journal") -> int:
    snap = sheet_snapshot(sheet_id, evaluate=False)
    rows = [
        resolve_formula_row(r, snap["columns"])
        for r in snap["rows"]
        if str(r.get("Account") or "").strip()
    ]
    if not rows:
        raise ValueError("journal sheet has no account rows")
    entry_date = str(rows[0].get("Date") or date.today().isoformat())
    lines = [
        {
            "account_code": row.get("Account"),
            "description": row.get("Description") or "",
            "debit": row.get("Debit") or 0,
            "credit": row.get("Credit") or 0,
            "employee_ref": row.get("Employee Ref") or "",
            "job_ref": row.get("Job Ref") or "",
        }
        for row in rows
    ]
    journal_id = post_journal_entry(
        entry_date,
        memo,
        lines,
        actor=actor,
        source_module="sheets",
        source_entity_type="sheet_book",
        source_entity_id=str(sheet_id),
    )
    with db() as con:
        con.execute(
            """UPDATE sheet_books SET status='posted',posted_entity_type='finance_journal',
               posted_entity_id=?,updated_at=CURRENT_TIMESTAMP WHERE id=?""",
            (str(journal_id), int(sheet_id)),
        )
    return journal_id
